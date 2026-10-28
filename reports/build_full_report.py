"""
LabPulse Academic Project Report Generator
CSE494: Intelligent NoSQL Databases (Academic Session 2026-27)

Generates:
1. reports/LabPulse_Project_Report.docx (Editable Microsoft Word document)
2. reports/LabPulse_Project_Report.pdf  (High-fidelity formatted PDF)
"""

import os
import sys
import json
import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether, HRFlowable
)
from reportlab.pdfgen import canvas

# ─── Directories and Inputs ──────────────────────────────────────────────────
BASE_DIR = r"d:\MongoProject"
REPORTS_DIR = os.path.join(BASE_DIR, "reports")
IMAGES_DIR = os.path.join(REPORTS_DIR, "images")
METRICS_JSON_PATH = os.path.join(BASE_DIR, "backend", "reports", "ml_evaluation_results.json")

os.makedirs(REPORTS_DIR, exist_ok=True)
os.makedirs(IMAGES_DIR, exist_ok=True)

with open(METRICS_JSON_PATH, "r", encoding="utf-8") as f:
    eval_data = json.load(f)

cls_m = eval_data["classification"]
reg_m = eval_data["regression"]
cm = cls_m["confusion_matrix"]

DOCX_OUTPUT_PATH = os.path.join(REPORTS_DIR, "LabPulse_Project_Report.docx")
PDF_OUTPUT_PATH = os.path.join(REPORTS_DIR, "LabPulse_Project_Report.pdf")

# ─── Report Content Definitions ─────────────────────────────────────────────

TITLE = "LabPulse: A MongoDB-Based Intelligent System for Laboratory Equipment Failure Prediction and Maintenance Management"
COURSE = "CSE494 — Intelligent NoSQL Databases"
SESSION = "Academic Session 2026–27"
PLACEHOLDER_STUDENT = "[Student Name]"
PLACEHOLDER_ROLL = "[Roll Number]"
PLACEHOLDER_DEPT = "Department of Computer Science and Engineering"
PLACEHOLDER_INST = "[Institution / University Name]"
PLACEHOLDER_GUIDE = "[Faculty Guide Name]"
SUBMISSION_DATE = "October 2026"

SYLLABUS_MAP = [
    ("PyMongo Connection & Configuration", "MongoClient connection pooling, URI configuration, server ping verification, and graceful fallback handling in backend/db.py.", "db.py"),
    ("Document Database & Schema Modeling", "Design of 5 specialized collections: equipment, faults, maintenance, predictions, model_metrics in labpulse_db.", "db.py, app.py"),
    ("Referencing vs. Embedding Tradeoff", "Embedded telemetry specs for 1:1 atomic unit access; referenced fault and maintenance documents for unbounded 1:N historical audits without 16MB document limit risks.", "db.py, models"),
    ("MongoDB JSON Schema Validation", "Strict $jsonSchema validation rules enforced on equipment, faults, and maintenance collections guaranteeing data integrity and type safety.", "db.py"),
    ("Single-Field, Compound & Text Indexes", "Single-field (equipment_id, status), compound (lab_id, status; equipment_id, scheduled_date), and full-text indexes ($text on name, ID, specs.cpu).", "db.py, insights"),
    ("MongoDB CRUD Operations", "Full transactional REST endpoints for equipment registration, real-time telemetry updates, fault incident logging, and maintenance scheduling.", "app.py"),
    ("MongoDB Aggregation Framework", "Multi-stage aggregation pipelines ($match, $group, $project, $sort, $avg, $facet) for KPI metrics, failure distributions, and lab utilization stats.", "app.py (/api/insights)"),
    ("Query Optimization & explain()", "Diagnostic query inspection with explain('executionStats'), analyzing IXSCAN vs. COLLSCAN index coverage and scan-to-return ratios.", "app.py, Analytics.jsx"),
    ("Predictive ML Integration with PyMongo", "Live extraction of 12-dimensional feature vectors from MongoDB operational state to feed trained Random Forest classification and regression models.", "ml/predictor.py, app.py"),
    ("End-to-End Functional Mini-Project", "Complete 4-tier mission-control dashboard with Three.js 3D computer lab digital twin, reactive UI, Flask API, and scikit-learn predictive engine.", "Entire Repository")
]

FIGURES = [
    ("Figure 3.1", "architecture_diagram.png", "LabPulse 4-Tier Architectural Pipeline and Component Interaction Diagram"),
    ("Figure 9.1", "fig_9_1_dashboard_overview.png", "LabPulse Overview Dashboard with Three.js 3D Computer Lab and Operational Health Counters"),
    ("Figure 9.2", "fig_9_2_3d_hover_telemetry.png", "Interactive 3D Workstation Raycasting with Real-Time Telemetry Tooltip and Status Badging"),
    ("Figure 9.3", "fig_9_3_equipment_registry.png", "Centralized Equipment Registry with Search, Status Badges, and Workstation Telemetry Cards"),
    ("Figure 9.4", "fig_9_4_fault_logs_reporting.png", "Fault Incident Logging Interface with Severity Indicators, Status Filters, and Repair Tracking"),
    ("Figure 9.5", "fig_9_5_maintenance_schedule.png", "Preventive Maintenance Scheduling System with Workstation Selector and Task Calendar"),
    ("Figure 9.6", "fig_9_6_database_insights.png", "MongoDB Database Insights Panel displaying Active JSON Schema Validators and Index Catalog"),
    ("Figure 9.7", "fig_9_7_predictions_dashboard.png", "Live Predictions Dashboard showcasing 7-Day Failure Risk Tiers and Conditional Repair Estimates"),
    ("Figure 9.8", "fig_9_8_model_evaluation_metrics.png", "Model Evaluation Metrics Panel showing Classification Scores, Confusion Matrix, and R² Curve"),
    ("Figure 9.9", "fig_9_9_prediction_drawer.png", "Equipment Detail Drawer revealing Live Feature Vectors and Explainable Risk Factors"),
    ("Figure 9.10", "fig_9_10_index_definitions.png", "Expanded Index Definitions and Query Optimization Statistics (IXSCAN execution details)")
]

# ─────────────────────────────────────────────────────────────────────────────
# DOCX GENERATION FUNCTION
# ─────────────────────────────────────────────────────────────────────────────

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_color="F1F5F9", border_color="0284C7"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"NOTE: {title}\n")
    run_t.bold = True
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(15, 23, 42)
    
    run_b = p.add_run(text)
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(51, 65, 85)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_docx_report():
    print("Building Word Report (.docx)...")
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header / Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_ft.text = f"LabPulse — CSE494 Academic Project Report | Page "
        p_ft.runs[0].font.size = Pt(8.5)
        p_ft.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # ── Title Page ───────────────────────────────────────────────────────────
    for _ in range(2):
        doc.add_paragraph()
    
    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_course = p_course.add_run(f"{COURSE.upper()}\n{SESSION.upper()}\nPROJECT REPORT")
    r_course.font.size = Pt(12)
    r_course.font.bold = True
    r_course.font.color.rgb = RGBColor(2, 132, 199)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(24)
    r_title = p_title.add_run("LABPULSE\n")
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    
    r_sub = p_title.add_run("A MongoDB-Based Intelligent System for Laboratory Equipment Failure Prediction and Maintenance Management")
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph()
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(36)
    p_meta.paragraph_format.line_spacing = 1.3
    
    r_subm = p_meta.add_run("Submitted in partial fulfillment of the requirements for the award of the degree of\n")
    r_subm.font.size = Pt(11)
    r_subm.font.color.rgb = RGBColor(100, 116, 139)
    
    r_deg = p_meta.add_run("Bachelor of Technology in Computer Science and Engineering\n\n")
    r_deg.font.size = Pt(12)
    r_deg.font.bold = True
    r_deg.font.color.rgb = RGBColor(15, 23, 42)
    
    tbl_author = doc.add_table(rows=2, cols=2)
    tbl_author.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_author.autofit = False
    tbl_author.columns[0].width = Inches(3.2)
    tbl_author.columns[1].width = Inches(3.2)
    
    c00 = tbl_author.cell(0, 0)
    c00.paragraphs[0].text = "SUBMITTED BY:"
    c00.paragraphs[0].runs[0].font.bold = True
    c00.paragraphs[0].runs[0].font.size = Pt(10)
    c00.paragraphs[0].runs[0].font.color.rgb = RGBColor(100, 116, 139)
    p00 = c00.add_paragraph(f"{PLACEHOLDER_STUDENT}\nRoll No: {PLACEHOLDER_ROLL}")
    p00.runs[0].font.bold = True
    p00.runs[0].font.size = Pt(11)
    
    c01 = tbl_author.cell(0, 1)
    c01.paragraphs[0].text = "UNDER THE GUIDANCE OF:"
    c01.paragraphs[0].runs[0].font.bold = True
    c01.paragraphs[0].runs[0].font.size = Pt(10)
    c01.paragraphs[0].runs[0].font.color.rgb = RGBColor(100, 116, 139)
    p01 = c01.add_paragraph(f"{PLACEHOLDER_GUIDE}\nDepartment of CSE")
    p01.runs[0].font.bold = True
    p01.runs[0].font.size = Pt(11)
    
    c10 = tbl_author.cell(1, 0)
    c10.paragraphs[0].text = f"\n{PLACEHOLDER_DEPT}\n{PLACEHOLDER_INST}"
    c10.paragraphs[0].runs[0].font.size = Pt(10.5)
    c10.paragraphs[0].runs[0].font.color.rgb = RGBColor(71, 85, 105)
    
    c11 = tbl_author.cell(1, 1)
    c11.paragraphs[0].text = f"\nDate of Submission: {SUBMISSION_DATE}\nAcademic Session: 2026–2027"
    c11.paragraphs[0].runs[0].font.size = Pt(10.5)
    c11.paragraphs[0].runs[0].font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # ── Executive Summary ────────────────────────────────────────────────────
    h_abs = doc.add_heading("Executive Summary", level=1)
    h_abs.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_paragraph(
        "Modern university computing laboratories represent mission-critical infrastructure supporting hundreds of students, "
        "faculty researchers, and technical staff across academic semesters. However, traditional computer laboratory management "
        "is predominantly reactive: hardware failures, thermal bottlenecks, and system crashes are addressed only after student reports "
        "or mid-session class disruptions occur. This operational inefficiency leads to cascading downtime, uncoordinated maintenance schedules, "
        "and lost academic productivity."
    )
    doc.add_paragraph(
        "LabPulse is an intelligent, full-stack, enterprise-grade laboratory equipment monitoring and predictive maintenance platform "
        "developed specifically to apply the theoretical principles and practical paradigms of CSE494: Intelligent NoSQL Databases. "
        "By synthesizing MongoDB 7.x document databases, PyMongo driver interfaces, Python Flask RESTful micro-services, scikit-learn machine learning, "
        "and a reactive Three.js 3D digital twin dashboard, LabPulse demonstrates the end-to-end transformation of raw database telemetry into "
        "actionable operational intelligence."
    )
    doc.add_paragraph(
        "The project rigorously satisfies all syllabus requirements of CSE494: rich document schema modeling with hybrid referencing and embedding; "
        "formal database schema validation via $jsonSchema; multi-tiered indexing strategies (single-field, compound, and text search); "
        "high-performance aggregation pipelines ($match, $group, $project, $sort, $avg, $facet); transparent query plan diagnostics via explain('executionStats'); "
        "and production-grade predictive modeling utilizing Random Forest classification (88.33% accuracy, 0.8462 F1-score) and fault-conditioned regression "
        "(R² = 0.7054, MAE = 1.39 hours)."
    )
    
    add_callout(doc, "Academic Course Context & Scope", 
                "This academic project report documents the design, mathematical foundations, database schema modeling, aggregation engineering, "
                "machine learning workflows, and empirical validation results for CSE494 (Intelligent NoSQL Databases). All machine learning models "
                "were evaluated on reproducible synthetic demonstration scenarios (1,200 records, Seed 42), strictly segregated from real operational "
                "equipment records.")

    # ── Table of Contents ────────────────────────────────────────────────────
    doc.add_heading("Table of Contents", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    toc_items = [
        ("Syllabus-to-Implementation Mapping Matrix", "Page 3"),
        ("Chapter 1: Introduction", "Page 4"),
        ("Chapter 2: Project Objectives and Scope", "Page 6"),
        ("Chapter 3: Requirements and System Architecture", "Page 8"),
        ("Chapter 4: MongoDB Data Modeling and Implementation", "Page 11"),
        ("Chapter 5: MongoDB Aggregation and Query Optimization", "Page 15"),
        ("Chapter 6: Dataset Description and Preprocessing", "Page 19"),
        ("Chapter 7: Machine-Learning Model Implementation", "Page 22"),
        ("Chapter 8: Evaluation Metrics and Results", "Page 26"),
        ("Chapter 9: Application Screens and Workflows", "Page 30"),
        ("Chapter 10: Testing and Verification", "Page 38"),
        ("Chapter 11: Installation and Execution Instructions", "Page 41"),
        ("Chapter 12: Conclusion, Limitations and Future Scope", "Page 44"),
        ("References and Course Bibliography", "Page 46")
    ]
    
    tbl_toc = doc.add_table(rows=len(toc_items), cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False
    tbl_toc.columns[0].width = Inches(5.2)
    tbl_toc.columns[1].width = Inches(1.3)
    
    for i, (title_text, page_str) in enumerate(toc_items):
        c0 = tbl_toc.cell(i, 0)
        c1 = tbl_toc.cell(i, 1)
        c0.paragraphs[0].text = title_text
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c1.paragraphs[0].text = page_str
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_cell_margins(c0, top=40, bottom=40, left=60, right=60)
        set_cell_margins(c1, top=40, bottom=40, left=60, right=60)

    doc.add_page_break()

    # ── Syllabus Mapping Table ───────────────────────────────────────────────
    doc.add_heading("CSE494 Syllabus-to-Implementation Mapping Matrix", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph(
        "The following matrix establishes direct, verifiable alignment between the CSE494: Intelligent NoSQL Databases syllabus modules "
        "and the concrete architectural components, algorithms, and source files implemented in the LabPulse system."
    )
    
    tbl_syl = doc.add_table(rows=len(SYLLABUS_MAP) + 1, cols=3)
    tbl_syl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_syl.autofit = False
    tbl_syl.columns[0].width = Inches(2.0)
    tbl_syl.columns[1].width = Inches(3.2)
    tbl_syl.columns[2].width = Inches(1.3)
    
    # Headers
    headers = ["Syllabus Topic & Practical", "Implemented Functionality in LabPulse", "Key Source Artifacts"]
    for j, h in enumerate(headers):
        cell = tbl_syl.cell(0, j)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.text = h
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
    for i, (top, func, src) in enumerate(SYLLABUS_MAP):
        row_idx = i + 1
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate([top, func, src]):
            cell = tbl_syl.cell(row_idx, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.text = val
            p.runs[0].font.size = Pt(8.5)
            if j == 0:
                p.runs[0].font.bold = True

    doc.add_page_break()

    # ── Chapter 1: Introduction ──────────────────────────────────────────────
    doc.add_heading("Chapter 1 — Introduction", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("1.1 Academic Laboratory Operational Context", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Academic computer laboratories in university computer science and engineering departments serve as high-density computing hubs. "
        "A typical department operates multiple specialized laboratories—such as the Systems & Operating Systems Lab, Artificial Intelligence "
        "and High-Performance Computing Center, Networking & Security Lab, and Software Engineering Development Suites. Each laboratory houses "
        "dozens of computing nodes, edge servers, network switches, and uninterruptible power supplies (UPS). These nodes experience heavy multi-user "
        "workloads including compilation pipelines, virtualization, GPU-accelerated neural network training, and concurrent distributed database clusters."
    )
    
    doc.add_heading("1.2 Operational Challenges & The Need for Predictive Monitoring", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Managing this heterogeneous hardware infrastructure presents significant operational bottlenecks when relying on legacy manual approaches:\n"
        "1. Reactive Maintenance Cycles: Hardware failures (e.g., thermal throttling, RAM fault allocation, fan failures, PSU drops) are typically "
        "identified only after students attempt lab practicals, resulting in wasted classroom hours and emergency technician callouts.\n"
        "2. Lack of Centralized Telemetry: Vital signs such as temperature, CPU load, RAM utilization, and continuous uptime are locked inside isolated "
        "operating system counters, remaining invisible to laboratory supervisors.\n"
        "3. Siloed Maintenance & Fault Histories: Incident reports and preventive maintenance logs are traditionally recorded in physical paper registers "
        "or unindexed spreadsheets. This precludes longitudinal analysis of recurring equipment faults or component failure patterns.\n"
        "4. Absence of Predictive Intelligence: Technicians have no mechanism to anticipate which workstations are at imminent risk of repeating a failure "
        "within upcoming critical periods (e.g., examination weeks), or how many labor hours a repair is likely to require."
    )

    doc.add_heading("1.3 Introduction to the LabPulse Solution", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "LabPulse is an intelligent, NoSQL-powered laboratory equipment monitoring and predictive maintenance system engineered to address these challenges. "
        "Built from the ground up to embody the principles of CSE494: Intelligent NoSQL Databases, LabPulse integrates:\n"
        "• A high-performance MongoDB document database capturing rich, hierarchical equipment telemetry, fault logs, and preventive maintenance schedules.\n"
        "• Strict schema validation ($jsonSchema) ensuring zero data corruption across multi-user environments.\n"
        "• Advanced aggregation pipelines and single/compound/text indexes optimizing real-time analytics.\n"
        "• A scikit-learn machine learning engine featuring a Random Forest Classifier to estimate 7-day repeat failure risk and a Random Forest Regressor "
        "to estimate required repair duration hours.\n"
        "• A cutting-edge React + Three.js interactive 3D digital twin dashboard providing real-time spatial awareness of laboratory workstations."
    )

    # ── Chapter 2: Project Objectives and Scope ──────────────────────────────
    doc.add_heading("Chapter 2 — Project Objectives and Scope", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("2.1 Primary Objectives", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The primary engineering and educational objectives of the LabPulse project are:\n"
        "1. Equipment Lifecycle Management: Provide an intuitive, real-time registry for cataloging computer laboratory workstations, servers, and network equipment.\n"
        "2. Fault Incident Tracking: Establish a rigorous incident tracking workflow from initial fault reporting, status progression ('reported' → 'in_progress' → 'resolved'), "
        "through root-cause documentation and resolution timestamps.\n"
        "3. Preventive Maintenance Scheduling: Implement a structured maintenance calendar tracking upcoming tasks, maintenance categories, and historical execution records.\n"
        "4. Deep MongoDB Integration: Fully utilize MongoDB as an intelligent document store—implementing validation rules, indexing architectures, aggregation pipelines, and explain() query plan optimization.\n"
        "5. Intelligent Failure Prediction: Develop, train, evaluate, and deploy machine learning models capable of predicting 7-day repeat failure probabilities and estimating repair durations.\n"
        "6. Unified Mission-Control Experience: Deliver a highly responsive single-page application (SPA) featuring spatial 3D visualization, interactive telemetry drawers, and real-time database diagnostics."
    )

    doc.add_heading("2.2 Implemented Scope", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The implemented scope encompasses six core operational views and backend subsystems:\n"
        "• Overview Dashboard: Featuring health metric scorecards, telemetry gauges, and an interactive Three.js 3D laboratory scene supporting orbit controls, raycast selection, and tooltip inspection.\n"
        "• Equipment Registry: A comprehensive data table with full-text search, lab filtering, health status badges, and rapid CRUD modal workflows.\n"
        "• Fault Logs: A specialized incident management interface facilitating fault ticket creation, status updates, technician assignment, and resolution recording.\n"
        "• Maintenance Scheduler: A task calendar managing upcoming preventive maintenance, recurring inspections, and completion status.\n"
        "• Analytics & Database Insights: Aggregation-driven KPI charts, collection metrics, interactive $jsonSchema inspector, index catalog, and explain('executionStats') query diagnostics.\n"
        "• Predictions Dashboard: Real-time risk scoring, 3-tier risk classification, conditional repair duration estimation, and interactive model performance evaluation tabs."
    )

    doc.add_heading("2.3 Technical Scope Boundaries & Limitations", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "To maintain academic integrity and project feasibility, the following boundaries are explicitly delineated:\n"
        "• Synthetic Training Data: The machine learning models are trained on a controlled, reproducible synthetic dataset of 1,200 scenarios (Seed 42). "
        "Real-world predictive deployment would require multi-semester historical telemetry collected from physical IoT sensors.\n"
        "• Advisory Predictive Role: Predictions generated by the system are advisory decision-support signals designed to assist human lab administrators in prioritizing preventive inspections. "
        "They do not constitute deterministic guarantees of hardware failure or precise repair times."
    )

    # ── Chapter 3: Requirements and System Architecture ──────────────────────
    doc.add_heading("Chapter 3 — Requirements and System Architecture", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("3.1 Technology Stack", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "LabPulse is architected using a modern, multi-tiered technology stack:\n"
        "• Client Presentation Tier: React 18, Vite 5.x build tool, Tailwind CSS for reactive styling, Three.js (r128) for 3D WebGL rendering, and Recharts for dynamic telemetry data visualization.\n"
        "• Application / REST API Tier: Python 3.11 with Flask 3.x, Flask-CORS for cross-origin resource sharing, python-dotenv for environment management, and Werkzeug utilities.\n"
        "• Database & Driver Tier: MongoDB Community Server 7.x running locally on mongodb://localhost:27017, communicated via PyMongo 4.x driver.\n"
        "• Machine Learning Tier: scikit-learn 1.4+, pandas 2.x, NumPy 1.26+, and joblib 1.3+ for model training, preprocessing pipeline assembly, and serialized artifact storage."
    )

    doc.add_heading("3.2 4-Tier Architectural Pipeline", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The LabPulse system is structured into four distinct, loosely coupled architectural tiers:\n"
        "1. Presentation Tier: Single-Page Application (SPA) providing mission-control navigation, 3D digital twin interaction, equipment telemetry cards, and live prediction tables.\n"
        "2. Application Tier: RESTful micro-service layer exposing endpoints for CRUD operations, aggregation pipelines, explain() query performance testing, and live model inference.\n"
        "3. Database Tier: MongoDB database named labpulse_db maintaining 5 collections with strict $jsonSchema validation rules and multi-tiered index coverage.\n"
        "4. Machine Learning Engine: Dual-pipeline architecture maintaining an offline synthetic training routine that serializes preprocessors and estimators, "
        "paired with an online real-time inference coordinator that extracts live MongoDB operational states and computes instantaneous predictions."
    )

    # Insert Architecture Diagram
    arch_img_path = os.path.join(IMAGES_DIR, "architecture_diagram.png")
    if os.path.exists(arch_img_path):
        doc.add_paragraph().paragraph_format.space_before = Pt(8)
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(arch_img_path, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 3.1: LabPulse 4-Tier Architectural Pipeline and Component Interaction Diagram")
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(71, 85, 105)
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── Chapter 4: MongoDB Data Modeling and Implementation ──────────────────
    doc.add_heading("Chapter 4 — MongoDB Data Modeling and Implementation", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("4.1 Database Architecture and Collection Topology", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "All application data is hosted within the database labpulse_db across five specialized collections:\n"
        "1. equipment (24 documents): Master hardware catalog containing workstation IDs, human-readable names, laboratory assignments, hardware specifications, operational statuses, and live telemetry.\n"
        "2. faults (18 documents): Incident tickets recording failure descriptions, categories, severity ratings, reporting dates, technician notes, and resolution timestamps.\n"
        "3. maintenance (16 documents): Scheduled and historical preventive maintenance tasks, maintenance types, schedules, and completion logs.\n"
        "4. predictions (24 documents): Cached machine learning inference records linking equipment IDs with predicted failure risk, confidence probabilities, and repair hours.\n"
        "5. model_metrics (1 document): Stored model evaluation metadata, confusion matrices, and training parameters enabling real-time dashboard auditability without retraining."
    )

    doc.add_heading("4.2 Referencing versus Embedding Architectural Tradeoffs", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "A foundational design decision in NoSQL database engineering is choosing between embedding subdocuments versus referencing normalized documents via foreign identifiers. "
        "In LabPulse, this decision follows established MongoDB best practices (Banker et al., 2016):\n"
        "• Embedded Design for Hardware Specifications: Subdocuments such as specs (cpu, ram_gb, storage_gb, gpu) and telemetry (temperature_c, cpu_usage_pct, ram_usage_pct) "
        "are directly embedded within the parent equipment document. These attributes exhibit a strict 1:1 relationship with the physical workstation, are always retrieved simultaneously "
        "during inventory scans, and do not grow unboundedly.\n"
        "• Referenced Design for Faults and Maintenance: Fault incidents and preventive maintenance logs are stored in dedicated collections and reference the parent workstation via equipment_id. "
        "This 1:N relationship is unbounded—a high-traffic workstation may accumulate dozens of fault tickets and maintenance visits across multiple academic years. Embedding these unbounded arrays "
        "inside the equipment document would lead to document bloat, risk exceeding MongoDB's 16MB BSON document limit, and cause severe write amplification during routine telemetry updates."
    )

    doc.add_heading("4.3 MongoDB JSON Schema Validation ($jsonSchema)", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "To ensure strict data integrity across API endpoints, LabPulse configures MongoDB JSON Schema validation ($jsonSchema) on the core collections:\n"
        "• equipment collection validator enforces: bsonType='object', required=['equipment_id', 'name', 'type', 'lab_id', 'status'], with status constrained to enum: ['operational', 'degraded', 'offline', 'maintenance'].\n"
        "• faults collection validator enforces: required=['fault_id', 'equipment_id', 'severity', 'status'], with severity constrained to ['low', 'medium', 'high', 'critical'] and status to ['reported', 'in_progress', 'resolved'].\n"
        "• maintenance collection validator enforces: required=['maintenance_id', 'equipment_id', 'type', 'status', 'scheduled_date'], with status constrained to ['scheduled', 'in_progress', 'completed', 'cancelled']."
    )

    doc.add_heading("4.4 Indexing Strategy and Implementation", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "To support high-throughput lookups and eliminate full collection scans (COLLSCAN), the following indexes are actively configured:\n"
        "1. Single-Field Unique Indexes: equipment_id on equipment, fault_id on faults, maintenance_id on maintenance, and equipment_id on predictions. These guarantee business key uniqueness.\n"
        "2. Single-Field Secondary Indexes: status and lab_id on equipment; equipment_id and status on faults; scheduled_date on maintenance.\n"
        "3. Compound Indexes: (lab_id: 1, status: 1) on equipment, optimizing simultaneous laboratory filtering and health state sorting; (equipment_id: 1, status: 1) on faults.\n"
        "4. Full-Text Index: Text index on equipment covering name, equipment_id, type, and specs.cpu, enabling instant multi-term keyword search across the registry."
    )

    doc.add_heading("4.5 Safe Deletion Semantics & Historical Audit Preservation", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "In production laboratory environments, deleting an equipment record must never cause catastrophic loss of historical fault tickets or maintenance compliance audits. "
        "LabPulse implements safe deletion semantics: when a workstation is decommissioned via the API, the system verifies its active fault commitments. The historical fault and maintenance records "
        "remain immutably preserved in their respective collections with equipment_id retained for institutional accreditation audits."
    )

    # ── Chapter 5: MongoDB Aggregation and Query Optimization ────────────────
    doc.add_heading("Chapter 5 — MongoDB Aggregation and Query Optimization", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("5.1 Analytical Aggregation Pipelines", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "MongoDB's Aggregation Framework provides a native, multi-stage data processing pipeline that processes documents on the database server without incurring network serialization overhead. "
        "LabPulse implements several key analytical pipelines exposed via /api/insights/aggregation:\n"
        "1. Equipment Status Breakdown: Uses $group by '$status' with '$count' accumulator to generate real-time fleet health ratios.\n"
        "2. Workstation Workload by Laboratory: Aggregates across lab_id using $group with $avg on '$telemetry.cpu_usage_pct', '$telemetry.ram_usage_pct', and '$telemetry.temperature_c'.\n"
        "3. Average Repair Duration by Category: Matches resolved faults via $match: { status: 'resolved' }, computes elapsed repair hours using $subtract on date fields, "
        "and aggregates with $group: { _id: '$category', avg_repair_hours: { $avg: '$repair_hours' } } followed by $sort.\n"
        "4. Multi-Faceted System Dashboard ($facet): Combines equipment health distributions, open fault severities, and upcoming maintenance workloads in a single round-trip database query."
    )

    doc.add_heading("5.2 Query Execution Diagnostics with explain('executionStats')", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Understanding query execution plans is critical for database optimization (Phaltankar et al., 2020). LabPulse provides real-time query performance diagnostics "
        "using MongoDB's cursor.explain('executionStats') method, exposed directly in the Database Insights panel:\n"
        "• Index Scan (IXSCAN): Lookups filtered by equipment_id or (lab_id, status) utilize the B-tree index, examining only the matching index keys (totalKeysExamined = nReturned) "
        "with stage FETCH, resulting in execution times below 1ms.\n"
        "• Collection Scan (COLLSCAN): In unindexed queries, MongoDB must inspect every document in the collection (totalDocsExamined = collection.count()), resulting in O(N) complexity.\n"
        "• Text Search (TEXT): Queries utilizing $text execute via the specialized text index, evaluating term frequencies and assigning textScore metadata."
    )

    add_callout(doc, "Query Latency on Academic Datasets", 
                "In academic environments with small demonstration datasets (e.g., 24 equipment records), elapsed query execution time differences between "
                "IXSCAN and COLLSCAN often measure in fractions of a millisecond. The true performance benefit of indexing is proven by comparing "
                "totalKeysExamined and totalDocsExamined against nReturned, demonstrating O(log N) scalability as data volumes expand.")

    # ── Chapter 6: Dataset Description and Preprocessing ────────────────────
    doc.add_heading("Chapter 6 — Dataset Description and Preprocessing", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("6.1 Synthetic Demonstration Dataset Architecture", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "To enable reproducible machine learning model training without exposing private institutional data, LabPulse generates a comprehensive synthetic dataset: "
        "backend/data/labpulse_ml_training.csv. The dataset comprises 1,200 simulated laboratory operational scenarios generated using a fixed random seed (Seed 42)."
    )

    doc.add_heading("6.2 Feature Space and Target Variables", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The feature space encompasses 12 carefully engineered operational variables representing physical telemetry, historical maintenance records, and equipment metadata:\n"
        "• Numerical Features (8): cpu_usage (%), ram_usage (%), temperature (°C), uptime (hours), previous_fault_count, unresolved_fault_count, "
        "days_since_last_maintenance, avg_historical_repair_hours.\n"
        "• Categorical Features (4): equipment_type ('workstation', 'server', 'network_switch'), lab_id ('LAB_A', 'LAB_B', 'LAB_C'), "
        "fault_category ('hardware', 'software', 'network', 'power', 'thermal', 'none'), fault_severity ('low', 'medium', 'high', 'critical', 'none').\n"
        "• Target Variables:\n"
        "  1. Classification Target: repeat_failure_within_7_days (binary: 0 or 1, indicating whether the workstation will fail again within one week).\n"
        "  2. Regression Target: repair_duration_hours (continuous: elapsed labor hours required to resolve an active fault, present on 548 active fault records)."
    )

    doc.add_heading("6.3 Preprocessing and Leakage Prevention Pipeline", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Data preprocessing is packaged into a scikit-learn ColumnTransformer:\n"
        "• Numerical attributes undergo median imputation (SimpleImputer(strategy='median')) to handle missing telemetry readings robustly against outliers.\n"
        "• Categorical attributes undergo most-frequent imputation followed by OneHotEncoder(handle_unknown='ignore') to prevent failures during novel category encounters.\n"
        "• Train-Test Partitioning: The dataset is split into an 80% training set (960 records) and a 20% held-out test set (240 records). Crucially, the preprocessing transformers "
        "are fitted strictly on the training partition and applied transform-only to the test partition, preventing data leakage."
    )

    # ── Chapter 7: Machine-Learning Model Implementation ────────────────────
    doc.add_heading("Chapter 7 — Machine-Learning Model Implementation", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("7.1 Repeat-Failure Classification Pipeline", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The repeat-failure classification model employs a Random Forest Classifier (RandomForestClassifier(n_estimators=100, max_depth=8, random_state=42)). "
        "Random Forest was selected due to its exceptional resistance to overfitting on tabular engineering telemetry, ability to capture non-linear interactions between temperature and load, "
        "and native capability to output well-calibrated class probabilities via predict_proba()."
    )

    doc.add_heading("7.2 Repair-Duration Regression Pipeline", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The repair duration estimation model employs a Random Forest Regressor (RandomForestRegressor(n_estimators=100, max_depth=8, random_state=42)). "
        "Trained on the 548 active fault records, the regressor models repair hours as a function of hardware severity, fault category, historical repair times, and equipment type."
    )

    doc.add_heading("7.3 Fault-Conditioned Prediction Rule", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "A critical engineering rule enforced by LabPulse is the Fault-Conditioned Prediction Rule: repair duration estimation is strictly computed only for equipment currently experiencing an active fault "
        "(status in ['degraded', 'offline'] with active fault severity and category). Workstations operating normally receive a null repair estimate ('Not applicable — no active fault'). "
        "This prevents nonsensical predictions such as predicting a 3-hour repair time for a completely healthy workstation."
    )

    doc.add_heading("7.4 Risk Tier Categorization & Thresholds", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The predicted failure probability is mapped into three operational risk tiers:\n"
        "• Low Risk (< 35% probability): Green indicator; workstation is operating within nominal parameters; standard routine maintenance schedule applies.\n"
        "• Moderate Risk (35% – 65% probability): Yellow indicator; telemetry displays mild thermal or load anomalies; technician advisory inspection recommended within 72 hours.\n"
        "• Elevated Risk (> 65% probability): Red indicator; high probability of repeat failure; immediate preventive diagnostic inspection prioritized."
    )

    # ── Chapter 8: Evaluation Metrics and Results ────────────────────────────
    doc.add_heading("Chapter 8 — Evaluation Metrics and Results", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph(
        "The machine learning models were comprehensively evaluated against the held-out test split (240 records for classification, 110 records for regression). "
        "The exact metrics retrieved directly from backend/reports/ml_evaluation_results.json are documented below."
    )

    # Classification Metrics Table
    tbl_cls = doc.add_table(rows=5, cols=3)
    tbl_cls.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cls.autofit = False
    tbl_cls.columns[0].width = Inches(2.2)
    tbl_cls.columns[1].width = Inches(1.3)
    tbl_cls.columns[2].width = Inches(3.0)
    
    cls_headers = ["Classification Metric", "Measured Value", "Engineering Interpretation"]
    for j, h in enumerate(cls_headers):
        cell = tbl_cls.cell(0, j)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        p = cell.paragraphs[0]
        p.text = h
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9.5)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
    cls_rows = [
        ("Accuracy", f"{cls_m['accuracy'] * 100:.2f}%", "Overall fraction of correct failure and non-failure classifications across the test split."),
        ("Precision", f"{cls_m['precision'] * 100:.2f}%", "Fraction of predicted failure alerts that were true failures (low false alert burden)."),
        ("Recall (Sensitivity)", f"{cls_m['recall'] * 100:.2f}%", "Fraction of actual repeat failures successfully detected (critical for preventing downtime)."),
        ("F1-Score", f"{cls_m['f1_score']:.4f}", "Harmonic mean of precision and recall, demonstrating balanced diagnostic capability.")
    ]
    for i, (m_name, m_val, m_desc) in enumerate(cls_rows):
        r_idx = i + 1
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for j, text in enumerate([m_name, m_val, m_desc]):
            cell = tbl_cls.cell(r_idx, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.text = text
            p.runs[0].font.size = Pt(9)
            if j == 0 or j == 1:
                p.runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    # Confusion Matrix Table
    doc.add_heading("Classification Confusion Matrix", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        f"The confusion matrix evaluated on 240 held-out test scenarios (155 true negative class 0, 85 true positive class 1) reveals:"
    )
    
    tbl_cm = doc.add_table(rows=3, cols=3)
    tbl_cm.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cm.autofit = False
    tbl_cm.columns[0].width = Inches(2.2)
    tbl_cm.columns[1].width = Inches(2.1)
    tbl_cm.columns[2].width = Inches(2.1)
    
    cm_headers = ["Actual \\ Predicted", "Predicted: No Failure (0)", "Predicted: Failure Risk (1)"]
    for j, h in enumerate(cm_headers):
        cell = tbl_cm.cell(0, j)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        p = cell.paragraphs[0]
        p.text = h
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        
    cm_data = [
        ("Actual: No Failure (0)", f"True Negative (TN): {cm['true_negatives']}", f"False Positive (FP): {cm['false_positives']}"),
        ("Actual: Failure Risk (1)", f"False Negative (FN): {cm['false_negatives']}", f"True Positive (TP): {cm['true_positives']}")
    ]
    for i, (r_title, c1_text, c2_text) in enumerate(cm_data):
        r_idx = i + 1
        for j, text in enumerate([r_title, c1_text, c2_text]):
            cell = tbl_cm.cell(r_idx, j)
            bg = "EFF6FF" if j > 0 else "F8FAFC"
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            p.text = text
            p.runs[0].font.size = Pt(9)
            if j == 0:
                p.runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    # Regression Metrics
    doc.add_heading("Regression Performance Metrics", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        f"The Random Forest Regressor evaluated across 110 active fault test scenarios achieved:\n"
        f"• Coefficient of Determination (R² Score): {reg_m['r2_score']:.4f} (accounting for over 70.5% of variance in repair duration).\n"
        f"• Mean Absolute Error (MAE): {reg_m['mae_hours']:.2f} hours (average deviation of approximately 1 hour and 23 minutes between predicted and actual repair duration)."
    )

    add_callout(doc, "Academic Transparency & Data Scope", 
                "These evaluation scores were obtained under controlled, reproducible conditions using a fixed random seed (Seed 42) on synthetic "
                "scenarios. For production deployment in university laboratories, longitudinal telemetry would be gathered across semesters "
                "to retrain models on empirical hardware degradation data.")

    # ── Chapter 9: Application Screens and Workflows ─────────────────────────
    doc.add_heading("Chapter 9 — Application Screens and Workflows", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph(
        "The LabPulse user interface delivers a cohesive mission-control environment. Authentic screenshots captured from the running application "
        "are presented below, illustrating user workflows across all six functional areas."
    )

    for fig_id, img_name, fig_desc in FIGURES[1:]:  # skip architecture diagram already included
        img_p = os.path.join(IMAGES_DIR, img_name)
        if os.path.exists(img_p):
            doc.add_paragraph().paragraph_format.space_before = Pt(6)
            p_i = doc.add_paragraph()
            p_i.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_i.add_run().add_picture(img_p, width=Inches(6.0))
            p_c = doc.add_paragraph()
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_c.paragraph_format.space_after = Pt(12)
            r_c = p_c.add_run(f"{fig_id}: {fig_desc}")
            r_c.font.size = Pt(9)
            r_c.font.bold = True
            r_c.font.color.rgb = RGBColor(71, 85, 105)

    # ── Chapter 10: Testing and Verification ─────────────────────────────────
    doc.add_heading("Chapter 10 — Testing and Verification", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph(
        "A rigorous, multi-layered verification strategy was executed across both backend and frontend layers to guarantee system reliability:\n"
        "1. Automated API Test Suite (backend/test_api.py): Comprises 6 distinct test suites validating: equipment CRUD operations, duplicate ID rejection, "
        "fault lifecycle synchronization, maintenance scheduling, schema validation rejection of invalid types, index presence, aggregation pipeline math, "
        "query explain() structure, ML model file availability, and prediction batch idempotency.\n"
        "2. Frontend Production Build Verification: Tested using Vite's production bundling (npm run build). The React 18 single-page application compiled cleanly "
        "with zero syntax errors or unresolved imports, outputting optimized static assets in 803ms.\n"
        "3. Model Artifact Verification: Confirmed that backend/models/classifier_pipeline.joblib, backend/models/regressor_pipeline.joblib, "
        "backend/models/model_metadata.json, and backend/reports/ml_evaluation_results.json exist, are readable, and parse successfully as JSON.\n"
        "4. Non-Destructive Integrity Protocol: All tests run during final submission preparation operated in non-destructive read/validate mode, "
        "preserving the active MongoDB labpulse_db records and collection state without data corruption."
    )

    # ── Chapter 11: Installation and Execution Instructions ─────────────────
    doc.add_heading("Chapter 11 — Installation and Execution Instructions", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph(
        "Follow these step-by-step instructions to reproduce and run the LabPulse platform locally on a Windows workstation:"
    )
    
    steps = [
        ("Step 1: Prerequisites Verification", "Ensure Node.js (v18.x or v20.x), Python (v3.10 or v3.11), and MongoDB Community Server (v7.x) are installed and available in the system PATH."),
        ("Step 2: MongoDB Service Startup", "Launch MongoDB locally by running: `net start MongoDB` or starting `mongod.exe` on default port 27017."),
        ("Step 3: Backend Environment Configuration", "Navigate to `backend/` and copy the configuration template: `copy .env.example .env`. Verify that `MONGODB_URI=mongodb://localhost:27017` and `MONGODB_DB_NAME=labpulse_db`."),
        ("Step 4: Python Virtual Environment & Dependencies", "Create and activate a virtual environment:\n`python -m venv venv`\n`venv\\Scripts\\activate`\n`pip install -r backend/requirements.txt`"),
        ("Step 5: Machine Learning Model Training (Optional)", "The repository includes pre-trained model artifacts in `backend/models/`. To retrain the pipelines on the synthetic dataset: `python backend/ml/train_models.py`"),
        ("Step 6: Start Flask REST API Backend", "Run the Flask backend server from the project root:\n`python backend/app.py`\nThe backend initializes on `http://localhost:5000` with active CORS headers."),
        ("Step 7: Frontend Dependency Installation & Startup", "Open a separate terminal in the project root:\n`npm install`\n`npm run dev`\nThe Vite development server starts at `http://localhost:5173`."),
        ("Step 8: Application Access & Testing", "Open a web browser and navigate to `http://localhost:5173`. Run the automated test suite in a separate terminal: `python backend/test_api.py`.")
    ]
    for s_title, s_desc in steps:
        p = doc.add_paragraph()
        r = p.add_run(f"• {s_title}: ")
        r.bold = True
        r.font.size = Pt(10)
        p.add_run(s_desc).font.size = Pt(9.5)

    # ── Chapter 12: Conclusion, Limitations and Future Scope ────────────────
    doc.add_heading("Chapter 12 — Conclusion, Limitations and Future Scope", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    doc.add_heading("12.1 Project Conclusion", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "LabPulse successfully demonstrates the integration of modern NoSQL database systems, RESTful micro-services, machine learning analytics, "
        "and interactive 3D WebGL user interfaces into an enterprise-ready laboratory management platform. By translating the theoretical tenets of "
        "CSE494: Intelligent NoSQL Databases into concrete software artifacts, the project proves how MongoDB's flexible document model, schema validation, "
        "rich indexing, and aggregation pipelines can power predictive operational systems."
    )

    doc.add_heading("12.2 Academic Limitations", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "In adherence to academic integrity principles, the project acknowledges several limitations:\n"
        "• Model training was conducted on synthetic demonstration scenarios (1,200 records, Seed 42). Real-world accuracy is contingent upon long-term empirical data collection.\n"
        "• The predictions are decision-support heuristics designed to assist technician prioritization; they cannot account for external physical anomalies such as sudden power surges or liquid spills.\n"
        "• Workstation telemetry in the current iteration is simulated via periodic background pulses rather than real-time hardware kernel agent daemon broadcasts."
    )

    doc.add_heading("12.3 Realistic Future Scope", level=2).runs[0].font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "Modest, achievable enhancements for future development include:\n"
        "1. Lightweight Client Daemon: Developing a cross-platform Python or Rust daemon running on client PCs to broadcast authentic CPU, RAM, and thermal metrics over MQTT.\n"
        "2. RFID / QR Code Equipment Tagging: Integrating mobile camera QR scanners for technicians to instantly view workstation prediction profiles upon physical inspection.\n"
        "3. Automated Helpdesk Webhooks: Adding Discord or Slack webhook notifications when a workstation transitions to the 'Elevated Risk' tier."
    )

    # ── References ───────────────────────────────────────────────────────────
    doc.add_heading("References and Course Bibliography", level=1).runs[0].font.color.rgb = RGBColor(15, 23, 42)
    
    references = [
        "Phaltankar, A., Ahsan, J., Harrison, M., & Nedov, L. (2020). MongoDB Fundamentals: A Hands-On Guide to Using MongoDB and Atlas in the Real World. Packt Publishing.",
        "Banker, K., Bakkum, P., Verch, S., Garrett, D., & Hawkins, T. (2016). MongoDB in Action: Covers MongoDB version 3.0 (2nd ed.). Manning Publications.",
        "MongoDB, Inc. (2026). MongoDB Documentation: Aggregation Pipeline Stages, Schema Validation, and Indexing. https://www.mongodb.com/docs/manual/",
        "PyMongo Documentation Team. (2026). PyMongo 4.x: Python Driver for MongoDB. https://pymongo.readthedocs.io/",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., ... & Duchesnay, E. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "React Team. (2026). React: The Library for Web and Native User Interfaces. https://react.dev/",
        "Cabello, R., et al. (2026). Three.js: JavaScript 3D WebGL Graphics Library. https://threejs.org/"
    ]
    for i, ref in enumerate(references):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.first_line_indent = Inches(-0.4)
        r_num = p.add_run(f"[{i+1}] ")
        r_num.bold = True
        r_num.font.size = Pt(9.5)
        p.add_run(ref).font.size = Pt(9.5)

    doc.save(DOCX_OUTPUT_PATH)
    print(f"Word report saved successfully to: {DOCX_OUTPUT_PATH}")

# ─────────────────────────────────────────────────────────────────────────────
# PDF GENERATION FUNCTION (ReportLab)
# ─────────────────────────────────────────────────────────────────────────────

class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to compute total page count and draw running header/footer."""
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        if self._pageNumber == 1:
            return  # Suppress headers/footers on title page
            
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header
        self.drawString(54, 750, "LabPulse: A MongoDB-Based System for Equipment Failure Prediction")
        self.drawRightString(558, 750, "CSE494 Project Report")
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.6)
        self.line(54, 744, 558, 744)
        
        # Footer
        self.line(54, 46, 558, 46)
        self.drawString(54, 34, "Session 2026–27 | Department of Computer Science and Engineering")
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 34, page_text)
        self.restoreState()

def build_pdf_report():
    print("Building PDF Report (.pdf)...")
    doc = SimpleDocTemplate(
        PDF_OUTPUT_PATH,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    style_title = ParagraphStyle(
        "CoverTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=24,
        leading=28,
        alignment=1,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=12
    )
    style_subtitle = ParagraphStyle(
        "CoverSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=12,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#475569"),
        spaceAfter=24
    )
    style_h1 = ParagraphStyle(
        "CustomH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A"),
        spaceBefore=14,
        spaceAfter=8,
        keepWithNext=True
    )
    style_h2 = ParagraphStyle(
        "CustomH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1E293B"),
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True
    )
    style_body = ParagraphStyle(
        "CustomBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6
    )
    style_caption = ParagraphStyle(
        "CustomCaption",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor("#475569"),
        spaceBefore=4,
        spaceAfter=10
    )
    style_tbl_hdr = ParagraphStyle(
        "TableHdr",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )
    style_tbl_cell = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#1E293B")
    )
    style_callout = ParagraphStyle(
        "Callout",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#1E293B")
    )

    story = []

    # ── Title Page ───────────────────────────────────────────────────────────
    story.append(Spacer(1, 40))
    story.append(Paragraph(f"<b>{COURSE.upper()}</b><br/>{SESSION.upper()}<br/><b>PROJECT REPORT</b>", ParagraphStyle(
        "CourseHdr", alignment=1, fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#0284C7"), spaceAfter=24
    )))
    story.append(Paragraph("LABPULSE", style_title))
    story.append(Paragraph("A MongoDB-Based Intelligent System for Laboratory Equipment Failure Prediction and Maintenance Management", style_subtitle))
    
    story.append(HRFlowable(width="60%", thickness=1.5, color=colors.HexColor("#0284C7"), spaceBefore=10, spaceAfter=30))
    
    subm_text = (
        f"<b>Submitted in partial fulfillment of the requirements for the degree of</b><br/>"
        f"<b>Bachelor of Technology in Computer Science and Engineering</b><br/><br/><br/>"
        f"<b>Submitted By:</b><br/>{PLACEHOLDER_STUDENT}<br/>Roll Number: {PLACEHOLDER_ROLL}<br/><br/>"
        f"<b>Under the Guidance of:</b><br/>{PLACEHOLDER_GUIDE}<br/>Department of CSE<br/><br/><br/>"
        f"<b>{PLACEHOLDER_DEPT}</b><br/>{PLACEHOLDER_INST}<br/>Date of Submission: {SUBMISSION_DATE}"
    )
    story.append(Paragraph(subm_text, ParagraphStyle(
        "CoverDetails", alignment=1, fontName="Helvetica", fontSize=10, leading=15, textColor=colors.HexColor("#334155")
    )))
    story.append(PageBreak())

    # ── Executive Summary ────────────────────────────────────────────────────
    story.append(Paragraph("Executive Summary", style_h1))
    story.append(Paragraph(
        "Modern university computing laboratories represent mission-critical infrastructure supporting hundreds of students, "
        "faculty researchers, and technical staff across academic semesters. However, traditional computer laboratory management "
        "is predominantly reactive: hardware failures, thermal bottlenecks, and system crashes are addressed only after student reports "
        "or mid-session class disruptions occur. This operational inefficiency leads to cascading downtime, uncoordinated maintenance schedules, "
        "and lost academic productivity.", style_body
    ))
    story.append(Paragraph(
        "LabPulse is an intelligent, full-stack, enterprise-grade laboratory equipment monitoring and predictive maintenance platform "
        "developed specifically to apply the theoretical principles and practical paradigms of CSE494: Intelligent NoSQL Databases. "
        "By synthesizing MongoDB 7.x document databases, PyMongo driver interfaces, Python Flask RESTful micro-services, scikit-learn machine learning, "
        "and a reactive Three.js 3D digital twin dashboard, LabPulse demonstrates the end-to-end transformation of raw database telemetry into "
        "actionable operational intelligence.", style_body
    ))
    story.append(Paragraph(
        "The project rigorously satisfies all syllabus requirements of CSE494: rich document schema modeling with hybrid referencing and embedding; "
        "formal database schema validation via $jsonSchema; multi-tiered indexing strategies (single-field, compound, and text search); "
        "high-performance aggregation pipelines ($match, $group, $project, $sort, $avg, $facet); transparent query plan diagnostics via explain('executionStats'); "
        "and production-grade predictive modeling utilizing Random Forest classification (88.33% accuracy, 0.8462 F1-score) and fault-conditioned regression "
        "(R² = 0.7054, MAE = 1.39 hours).", style_body
    ))
    
    # Callout
    callout_data = [[Paragraph(
        "<b>NOTE: Academic Course Context & Scope</b><br/>"
        "This academic project report documents the design, mathematical foundations, database schema modeling, aggregation engineering, "
        "machine learning workflows, and empirical validation results for CSE494 (Intelligent NoSQL Databases). All machine learning models "
        "were evaluated on reproducible synthetic demonstration scenarios (1,200 records, Seed 42), strictly segregated from real operational "
        "equipment records.", style_callout
    )]]
    tbl_call = Table(callout_data, colWidths=[500])
    tbl_call.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F1F5F9")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('LINEBEFORE', (0,0), (0,-1), 3, colors.HexColor("#0284C7")),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
    ]))
    story.append(Spacer(1, 8))
    story.append(tbl_call)
    story.append(Spacer(1, 14))

    # ── Syllabus Mapping Table ───────────────────────────────────────────────
    story.append(Paragraph("CSE494 Syllabus-to-Implementation Mapping Matrix", style_h1))
    story.append(Paragraph("Direct alignment between syllabus modules and implemented codebase artifacts:", style_body))
    
    syl_table_data = [[
        Paragraph("<b>Syllabus Topic & Practical</b>", style_tbl_hdr),
        Paragraph("<b>Implemented Functionality in LabPulse</b>", style_tbl_hdr),
        Paragraph("<b>Key Source Artifacts</b>", style_tbl_hdr)
    ]]
    for top, func, src in SYLLABUS_MAP:
        syl_table_data.append([
            Paragraph(f"<b>{top}</b>", style_tbl_cell),
            Paragraph(func, style_tbl_cell),
            Paragraph(f"<code>{src}</code>", style_tbl_cell)
        ])
    tbl_syl_pdf = Table(syl_table_data, colWidths=[130, 260, 110])
    tbl_syl_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0F172A")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(tbl_syl_pdf)
    story.append(PageBreak())

    # ── Chapters 1 & 2 ───────────────────────────────────────────────────────
    story.append(Paragraph("Chapter 1 — Introduction", style_h1))
    story.append(Paragraph("1.1 Academic Laboratory Operational Context", style_h2))
    story.append(Paragraph(
        "Academic computer laboratories in university computer science and engineering departments serve as high-density computing hubs. "
        "A typical department operates multiple specialized laboratories—such as the Systems & Operating Systems Lab, Artificial Intelligence "
        "and High-Performance Computing Center, Networking & Security Lab, and Software Engineering Development Suites. Each laboratory houses "
        "dozens of computing nodes, edge servers, network switches, and uninterruptible power supplies (UPS). These nodes experience heavy multi-user "
        "workloads including compilation pipelines, virtualization, GPU-accelerated neural network training, and concurrent distributed database clusters.", style_body
    ))
    story.append(Paragraph("1.2 Operational Challenges & The Need for Predictive Monitoring", style_h2))
    story.append(Paragraph(
        "Managing this heterogeneous hardware infrastructure presents significant operational bottlenecks when relying on legacy manual approaches:\n"
        "1. Reactive Maintenance Cycles: Hardware failures (e.g., thermal throttling, RAM fault allocation, fan failures, PSU drops) are typically "
        "identified only after students attempt lab practicals, resulting in wasted classroom hours and emergency technician callouts.\n"
        "2. Lack of Centralized Telemetry: Vital signs such as temperature, CPU load, RAM utilization, and continuous uptime are locked inside isolated "
        "operating system counters, remaining invisible to laboratory supervisors.\n"
        "3. Siloed Maintenance & Fault Histories: Incident reports and preventive maintenance logs are traditionally recorded in physical paper registers "
        "or unindexed spreadsheets. This precludes longitudinal analysis of recurring equipment faults or component failure patterns.\n"
        "4. Absence of Predictive Intelligence: Technicians have no mechanism to anticipate which workstations are at imminent risk of repeating a failure "
        "within upcoming critical periods (e.g., examination weeks), or how many labor hours a repair is likely to require.", style_body
    ))

    story.append(Paragraph("Chapter 2 — Project Objectives and Scope", style_h1))
    story.append(Paragraph("2.1 Primary Objectives", style_h2))
    story.append(Paragraph(
        "The primary engineering and educational objectives of the LabPulse project are:\n"
        "1. Equipment Lifecycle Management: Provide an intuitive, real-time registry for cataloging computer laboratory workstations, servers, and network equipment.\n"
        "2. Fault Incident Tracking: Establish a rigorous incident tracking workflow from initial fault reporting, status progression ('reported' → 'in_progress' → 'resolved'), "
        "through root-cause documentation and resolution timestamps.\n"
        "3. Preventive Maintenance Scheduling: Implement a structured maintenance calendar tracking upcoming tasks, maintenance categories, and historical execution records.\n"
        "4. Deep MongoDB Integration: Fully utilize MongoDB as an intelligent document store—implementing validation rules, indexing architectures, aggregation pipelines, and explain() query plan optimization.\n"
        "5. Intelligent Failure Prediction: Develop, train, evaluate, and deploy machine learning models capable of predicting 7-day repeat failure probabilities and estimating repair durations.\n"
        "6. Unified Mission-Control Experience: Deliver a highly responsive single-page application (SPA) featuring spatial 3D visualization, interactive telemetry drawers, and real-time database diagnostics.", style_body
    ))

    # ── Chapter 3: Architecture & Diagram ────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("Chapter 3 — Requirements and System Architecture", style_h1))
    story.append(Paragraph("3.1 4-Tier Architectural Pipeline", style_h2))
    story.append(Paragraph(
        "The LabPulse system is structured into four distinct, loosely coupled architectural tiers:\n"
        "1. Presentation Tier: React 18 single-page application featuring Three.js 3D laboratory digital twin, Recharts telemetry curves, and Tailwind CSS.\n"
        "2. Application Tier: Python 3.11 with Flask RESTful API coordinating CRUD requests, aggregation pipelines, and inference calls.\n"
        "3. Database Tier: MongoDB 7.x database (labpulse_db) with PyMongo driver, JSON Schema validation, and multi-tier indexing.\n"
        "4. Machine Learning Engine: scikit-learn offline pipeline serializing Random Forest preprocessors and estimators, coupled with an online inference coordinator.", style_body
    ))

    arch_img_path = os.path.join(IMAGES_DIR, "architecture_diagram.png")
    if os.path.exists(arch_img_path):
        story.append(Spacer(1, 6))
        story.append(Image(arch_img_path, width=480, height=300))
        story.append(Paragraph("Figure 3.1: LabPulse 4-Tier Architectural Pipeline and Component Interaction Diagram", style_caption))

    # ── Chapter 4 & 5: Database & Aggregation ────────────────────────────────
    story.append(Paragraph("Chapter 4 — MongoDB Data Modeling and Implementation", style_h1))
    story.append(Paragraph("4.1 Referencing versus Embedding Architectural Tradeoffs", style_h2))
    story.append(Paragraph(
        "In LabPulse, data modeling strictly balances read performance with document boundary constraints (Banker et al., 2016):\n"
        "• Embedded Specs and Telemetry: Hardware specifications (specs) and operational sensor readings (telemetry) are directly embedded in equipment documents. "
        "Because these fields exhibit a 1:1 relationship with the physical node and are queried synchronously during dashboard renders, embedding maximizes atomic read throughput.\n"
        "• Referenced Faults and Maintenance: Incident tickets (faults) and preventive maintenance visits (maintenance) reference equipment_id. "
        "Because high-traffic workstations generate unbounded arrays of historical logs across academic years, embedding these would cause document bloat and risk exceeding "
        "MongoDB's 16MB BSON limit.", style_body
    ))
    story.append(Paragraph("4.2 MongoDB JSON Schema Validation ($jsonSchema)", style_h2))
    story.append(Paragraph(
        "To enforce data integrity across all operational endpoints, MongoDB $jsonSchema validators are applied to equipment, faults, and maintenance collections. "
        "The validators enforce BSON types, required keys, string regex patterns, and strict enum sets for status and severity.", style_body
    ))

    story.append(Paragraph("Chapter 5 — MongoDB Aggregation and Query Optimization", style_h1))
    story.append(Paragraph("5.1 Aggregation Pipeline Stages and Diagnostic Analysis", style_h2))
    story.append(Paragraph(
        "Analytical metrics are processed in-database via MongoDB Aggregation Pipelines ($match, $group, $project, $sort, $avg, $facet). "
        "Query execution performance is inspected via explain('executionStats'), verifying that index scans (IXSCAN) on indexed keys "
        "(e.g., equipment_id, (lab_id, status)) reduce totalKeysExamined to exact match counts, eliminating expensive collection scans (COLLSCAN).", style_body
    ))

    # ── Chapter 6, 7 & 8: Machine Learning & Results ─────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("Chapter 6 — Dataset Description and Preprocessing", style_h1))
    story.append(Paragraph(
        "The machine learning models are trained on backend/data/labpulse_ml_training.csv (1,200 simulated operational scenarios, Random Seed 42). "
        "Features include 8 numerical variables (cpu_usage, ram_usage, temperature, uptime, previous_fault_count, unresolved_fault_count, "
        "days_since_last_maintenance, avg_historical_repair_hours) and 4 categorical variables (equipment_type, lab_id, fault_category, fault_severity). "
        "Preprocessing uses ColumnTransformer with median imputation for numericals, and most-frequent imputation + OneHotEncoder for categoricals, "
        "strictly fitted on the 80% training split.", style_body
    ))

    story.append(Paragraph("Chapter 7 — Machine-Learning Model Implementation", style_h1))
    story.append(Paragraph(
        "The ML subsystem deploys a Random Forest Classifier for 7-day repeat failure risk prediction, alongside a Random Forest Regressor "
        "for repair duration hours estimation. Critically, the Fault-Conditioned Prediction Rule stipulates that repair hours are estimated "
        "only when equipment experiences an active fault; healthy workstations receive a null estimate ('Not applicable — no active fault').", style_body
    ))

    story.append(Paragraph("Chapter 8 — Evaluation Metrics and Results", style_h1))
    story.append(Paragraph(
        "The models were evaluated against the held-out test split (240 classification scenarios, 110 regression scenarios) retrieved directly from "
        "backend/reports/ml_evaluation_results.json:", style_body
    ))

    # Metrics Table
    m_table_data = [
        [Paragraph("<b>Metric</b>", style_tbl_hdr), Paragraph("<b>Score</b>", style_tbl_hdr), Paragraph("<b>Interpretation</b>", style_tbl_hdr)],
        [Paragraph("<b>Classifier Accuracy</b>", style_tbl_cell), Paragraph(f"<b>{cls_m['accuracy'] * 100:.2f}%</b>", style_tbl_cell), Paragraph("Overall correct failure and non-failure classifications across test partition.", style_tbl_cell)],
        [Paragraph("<b>Classifier Precision</b>", style_tbl_cell), Paragraph(f"<b>{cls_m['precision'] * 100:.2f}%</b>", style_tbl_cell), Paragraph("High precision guarantees minimal false alarms for lab technicians.", style_tbl_cell)],
        [Paragraph("<b>Classifier Recall</b>", style_tbl_cell), Paragraph(f"<b>{cls_m['recall'] * 100:.2f}%</b>", style_tbl_cell), Paragraph("90.59% sensitivity ensures critical hardware failures are not missed.", style_tbl_cell)],
        [Paragraph("<b>Classifier F1-Score</b>", style_tbl_cell), Paragraph(f"<b>{cls_m['f1_score']:.4f}</b>", style_tbl_cell), Paragraph("Harmonic mean confirms balanced diagnostic capability on imbalanced classes.", style_tbl_cell)],
        [Paragraph("<b>Confusion Matrix</b>", style_tbl_cell), Paragraph(f"<b>TN={cm['true_negatives']}, FP={cm['false_positives']}<br/>FN={cm['false_negatives']}, TP={cm['true_positives']}</b>", style_tbl_cell), Paragraph("Evaluated on 155 negative class (0) and 85 positive class (1) test samples.", style_tbl_cell)],
        [Paragraph("<b>Regressor R² Score</b>", style_tbl_cell), Paragraph(f"<b>{reg_m['r2_score']:.4f}</b>", style_tbl_cell), Paragraph("Model accounts for >70.5% of variance in active fault repair hours.", style_tbl_cell)],
        [Paragraph("<b>Regressor MAE</b>", style_tbl_cell), Paragraph(f"<b>{reg_m['mae_hours']:.2f} hours</b>", style_tbl_cell), Paragraph("Average repair duration error is approximately 1 hour and 23 minutes.", style_tbl_cell)]
    ]
    tbl_m_pdf = Table(m_table_data, colWidths=[130, 110, 260])
    tbl_m_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0F172A")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    story.append(tbl_m_pdf)
    story.append(Spacer(1, 10))

    # ── Chapter 9: Application Screens ───────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("Chapter 9 — Application Screens and Workflows", style_h1))
    story.append(Paragraph("Authentic screen captures from the running LabPulse application demonstrating full-stack workflows:", style_body))

    for fig_id, img_name, fig_desc in FIGURES[1:]:
        img_p = os.path.join(IMAGES_DIR, img_name)
        if os.path.exists(img_p):
            story.append(KeepTogether([
                Spacer(1, 4),
                Image(img_p, width=480, height=270),
                Paragraph(f"{fig_id}: {fig_desc}", style_caption),
                Spacer(1, 6)
            ]))

    # ── Chapter 10, 11, 12 & References ─────────────────────────────────────
    story.append(PageBreak())
    story.append(Paragraph("Chapter 10 — Testing and Verification", style_h1))
    story.append(Paragraph(
        "Automated backend testing (backend/test_api.py) verified 6 full suites: equipment CRUD, duplicate rejection, fault synchronization, "
        "maintenance tracking, schema validation rejection, index cataloging, aggregation pipelines, explain() executionStats, model serialization, "
        "and prediction API idempotency. The frontend production build (npm run build) completed in 803ms with zero errors.", style_body
    ))

    story.append(Paragraph("Chapter 11 — Installation and Execution Instructions", style_h1))
    story.append(Paragraph(
        "1. Start MongoDB: Ensure local MongoDB Community Server is running on port 27017.<br/>"
        "2. Backend Config: Copy <code>backend/.env.example</code> to <code>backend/.env</code>.<br/>"
        "3. Python Environment: Create venv and install: <code>pip install -r backend/requirements.txt</code>.<br/>"
        "4. Start Backend API: Run <code>python backend/app.py</code> (starts on http://localhost:5000).<br/>"
        "5. Start Frontend SPA: In project root, run <code>npm install</code> then <code>npm run dev</code> (starts on http://localhost:5173).<br/>"
        "6. Automated Verification: Run <code>python backend/test_api.py</code> to execute the full test suite.", style_body
    ))

    story.append(Paragraph("Chapter 12 — Conclusion, Limitations and Future Scope", style_h1))
    story.append(Paragraph(
        "LabPulse demonstrates an end-to-end intelligent NoSQL system adhering to all CSE494 syllabus practicals. "
        "Limitations include reliance on synthetic demonstration data (1,200 records, Seed 42) and the advisory nature of predictions. "
        "Future scope includes client OS telemetry daemons, mobile RFID/QR equipment scanning, and automated incident webhooks.", style_body
    ))

    story.append(Paragraph("References and Course Bibliography", style_h1))
    ref_text = (
        "[1] Phaltankar, A., Ahsan, J., Harrison, M., & Nedov, L. (2020). <i>MongoDB Fundamentals</i>. Packt Publishing.<br/>"
        "[2] Banker, K., Bakkum, P., Verch, S., Garrett, D., & Hawkins, T. (2016). <i>MongoDB in Action</i> (2nd ed.). Manning Publications.<br/>"
        "[3] MongoDB, Inc. (2026). <i>MongoDB Documentation: Aggregation Pipelines, Schema Validation & Indexing</i>.<br/>"
        "[4] PyMongo Documentation Team. (2026). <i>PyMongo: Python Driver for MongoDB Documentation</i>.<br/>"
        "[5] Pedregosa, F., et al. (2011). <i>Scikit-learn: Machine Learning in Python</i>. JMLR, 12, 2825-2830.<br/>"
        "[6] React Team. (2026). <i>React: Frontend SPA Framework</i>. https://react.dev/<br/>"
        "[7] Cabello, R., et al. (2026). <i>Three.js: WebGL 3D Computer Graphics Library</i>. https://threejs.org/"
    )
    story.append(Paragraph(ref_text, style_body))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF report saved successfully to: {PDF_OUTPUT_PATH}")

if __name__ == "__main__":
    build_docx_report()
    build_pdf_report()
    print("All Reports Built Successfully!")
